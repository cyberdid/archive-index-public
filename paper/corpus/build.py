"""Lay out the UA-PAR corpus from the raw sources (plan Task 4). Offline: collect.py fetches the raw data first.

The tree is out_dir/<project>/<file>, so `archive index --projects-root out_dir` indexes it like a personal archive.
File names come from stable ids (page id, dataset id, image id), never from titles: a title such as "Password
manager" would otherwise meet the archive's sensitive-name filter and silently vanish from the index.
manifest.jsonl lists every file once: doc_id (path relative to out_dir), part, lang, source, sha256, licence and the
fields each part needs later (title, transcript, sentence_id, attribution).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import random
import shutil
import subprocess
import tarfile
import tempfile
import warnings
from pathlib import Path

PARTS = ("notes", "documents", "tables", "code", "images", "speech")
TEXT_PARTS = ("notes", "documents", "tables")  # language shares are measured on these
DEFAULT_SIZE = {"notes": 3000, "documents": 400, "tables": 150, "code": 30, "images": 1000, "speech": 800}
CODE_MAX_FILE = 1024 * 1024


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _entry(out_dir: Path, path: Path, part: str, lang: str, source: str, licence: str, **extra) -> dict:
    return {"doc_id": str(path.relative_to(out_dir)), "path": str(path), "part": part, "lang": lang,
            "source": source, "sha256": _sha256(path), "licence": licence, **extra}


def _load_wiki(raw_dir: Path) -> list[dict]:
    pages = []
    for lang in ("uk", "ru"):
        for f in sorted((raw_dir / f"wiki_{lang}").glob("*.json")):
            page = json.loads(f.read_text(encoding="utf-8"))
            page["lang"] = lang
            pages.append(page)
    return pages


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in text.split("\n") if p.strip()]


def _write_note(page: dict, path: Path) -> None:
    body = "\n\n".join(_paragraphs(page["text"]))
    path.write_text(f"# {page['title']}\n\nДжерело: {page['url']}\n\n{body}\n", encoding="utf-8")


def _write_docx(page: dict, path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading(page["title"], level=1)
    for paragraph in _paragraphs(page["text"]):
        document.add_paragraph(paragraph)
    document.add_paragraph(f"Джерело: {page['url']}")
    document.save(str(path))


def _write_pptx(page: dict, path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    paragraphs = _paragraphs(page["text"])
    for start in range(0, max(len(paragraphs), 1), 3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = page["title"] if start == 0 else f"{page['title']} ({start // 3 + 1})"
        slide.placeholders[1].text = "\n".join(paragraphs[start:start + 3])[:1500]
    deck.save(str(path))


def _write_pdf(page: dict, path: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = Path(tmp) / (path.stem + ".docx")
        _write_docx(page, docx_path)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx_path)],
                       capture_output=True, check=True, timeout=180)
        shutil.move(str(Path(tmp) / (path.stem + ".pdf")), path)


def _layout_text(pages: list[dict], out_dir: Path, size: dict, render_pdf: bool, rng: random.Random) -> list[dict]:
    pool = sorted(pages, key=lambda p: (p["lang"], int(p["pageid"])))
    rng.shuffle(pool)
    docs, notes = pool[: size["documents"]], pool[size["documents"]: size["documents"] + size["notes"]]
    entries = []
    for page in notes:
        path = out_dir / page["topic"] / "notes" / f"{page['lang']}-{page['pageid']}.md"
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_note(page, path)
        entries.append(_entry(out_dir, path, "notes", page["lang"], page["url"], "CC-BY-SA-4.0", title=page["title"]))
    writers = [("docx", _write_docx), ("pptx", _write_pptx), ("pdf", _write_pdf)]
    for i, page in enumerate(docs):
        ext, writer = writers[i % 3]
        if ext == "pdf" and not render_pdf:
            ext, writer = writers[0]
        path = out_dir / page["topic"] / "docs" / f"{page['lang']}-{page['pageid']}.{ext}"
        path.parent.mkdir(parents=True, exist_ok=True)
        writer(page, path)
        entries.append(_entry(out_dir, path, "documents", page["lang"], page["url"], "CC-BY-SA-4.0",
                              title=page["title"]))
    return entries


def _layout_tables(raw_dir: Path, out_dir: Path, count: int) -> list[dict]:
    from openpyxl import Workbook

    entries = []
    for csv_path in sorted((raw_dir / "datagovua").glob("*.csv")):
        if len(entries) >= count:
            break
        raw = csv_path.read_bytes()
        # CKAN occasionally labels an XLSX or a one-line download link as CSV.
        if raw.startswith(b"PK\x03\x04"):
            warnings.warn(f"Skipping non-CSV resource: {csv_path.name}", RuntimeWarning)
            continue
        try:
            if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                text = raw.decode("utf-16")
            else:
                text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("cp1251")
            except UnicodeDecodeError:
                warnings.warn(f"Skipping undecodable CSV resource: {csv_path.name}", RuntimeWarning)
                continue
        lines = [line for line in text.splitlines() if line.strip()]
        if len(lines) < 2 or text.lstrip().lower().startswith(("<!doctype html", "<html")):
            warnings.warn(f"Skipping non-tabular CSV resource: {csv_path.name}", RuntimeWarning)
            continue
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            # Sniffer can fail on large/irregular Ukrainian datasets. Use the header's strongest separator;
            # when none exists, preserve a legitimate one-column CSV.
            header = lines[0]
            delimiter = max((",", ";", "\t"), key=header.count)
        rows = [row for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter) if row]
        if len(rows) < 2:
            warnings.warn(f"Skipping non-tabular CSV resource: {csv_path.name}", RuntimeWarning)
            continue
        meta = json.loads(csv_path.with_suffix(".json").read_text(encoding="utf-8"))
        folder = out_dir / "opendata" / "tables"
        folder.mkdir(parents=True, exist_ok=True)
        if len(entries) % 3 == 2:  # every third valid table as a spreadsheet, so both formats are present
            book = Workbook()
            sheet = book.active
            for row in rows:
                sheet.append(row)
            path = folder / f"{meta['id']}.xlsx"
            book.save(str(path))
        else:
            path = folder / f"{meta['id']}.csv"
            with path.open("w", encoding="utf-8", newline="") as fh:
                csv.writer(fh, delimiter=delimiter).writerows(rows)
        entries.append(_entry(out_dir, path, "tables", "uk", meta["url"], meta["licence"], title=meta.get("title", "")))
    if len(entries) < count:
        raise ValueError(f"Requested {count} valid open-data tables, found {len(entries)}")
    return entries


def _layout_code(raw_dir: Path, out_dir: Path, count: int) -> list[dict]:
    entries = []
    for repo in sorted(p for p in (raw_dir / "github").glob("*") if p.is_dir())[:count]:
        meta = json.loads((repo / "meta.json").read_text(encoding="utf-8"))
        target_root = out_dir / ("code-" + repo.name.split("__", 1)[-1].lower())
        for f in sorted(repo.rglob("*")):
            if not f.is_file() or f.name == "meta.json" or ".git" in f.parts or f.stat().st_size > CODE_MAX_FILE:
                continue
            target = target_root / f.relative_to(repo)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, target)
            entries.append(_entry(out_dir, target, "code", "mixed", meta["url"], meta["licence"]))
    return entries


def _layout_images(raw_dir: Path, out_dir: Path, count: int, rng: random.Random) -> list[dict]:
    attributions = {}
    with (raw_dir / "xm3600_attributions" / "image_attributions.csv").open(encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            attributions[row["ImageID"]] = row
    entries = []
    with tarfile.open(raw_dir / "xm3600_images" / "images.tgz") as tar:
        members = sorted((m for m in tar.getmembers() if m.isfile() and m.name.endswith(".jpg")), key=lambda m: m.name)
        for member in sorted(rng.sample(members, min(count, len(members))), key=lambda m: m.name):
            image_id = Path(member.name).stem
            path = out_dir / "photos" / f"{image_id}.jpg"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(tar.extractfile(member).read())
            row = attributions.get(image_id, {})
            entries.append(_entry(out_dir, path, "images", "-", row.get("OriginalLandingURL", ""),
                                  row.get("License", "CC-BY-2.0"), image_id=image_id, author=row.get("Author", "")))
    return entries


def _layout_speech(raw_dir: Path, out_dir: Path, count: int, rng: random.Random) -> list[dict]:
    entries = []
    per_language = count // 2
    for lang in ("uk", "ru"):
        rows = []
        with (raw_dir / f"fleurs_{lang}_text" / "test.tsv").open(encoding="utf-8") as fh:
            for cols in csv.reader(fh, delimiter="\t", quoting=csv.QUOTE_NONE):
                if len(cols) >= 3:
                    rows.append({"sentence_id": cols[0], "file": cols[1], "transcript": cols[2]})
        chosen = sorted(rng.sample(rows, min(per_language, len(rows))), key=lambda r: r["file"])
        wanted = {r["file"]: r for r in chosen}
        with tarfile.open(raw_dir / f"fleurs_{lang}_audio" / "test.tar.gz") as tar:
            for member in tar:
                name = Path(member.name).name
                if member.isfile() and name in wanted:
                    path = out_dir / "recordings" / lang / name
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_bytes(tar.extractfile(member).read())
                    r = wanted[name]
                    entries.append(_entry(out_dir, path, "speech", lang,
                                          "https://huggingface.co/datasets/google/fleurs", "CC-BY-4.0",
                                          transcript=r["transcript"], sentence_id=r["sentence_id"]))
    return entries


def build_corpus(raw_dir: Path, out_dir: Path, seed: int = 2026, size: dict[str, int] | None = None,
                 render_pdf: bool = True) -> list[dict]:
    size = {**DEFAULT_SIZE, **(size or {})}
    rng = random.Random(seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    entries = _layout_text(_load_wiki(raw_dir), out_dir, size, render_pdf, rng)
    entries += _layout_tables(raw_dir, out_dir, size["tables"])
    entries += _layout_code(raw_dir, out_dir, size["code"])
    entries += _layout_images(raw_dir, out_dir, size["images"], rng)
    entries += _layout_speech(raw_dir, out_dir, size["speech"], rng)
    with (out_dir / "manifest.jsonl").open("w", encoding="utf-8") as fh:
        for entry in entries:
            fh.write(json.dumps(entry, ensure_ascii=False) + "\n")
    return entries


def read_manifest(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def language_shares(entries: list[dict]) -> dict[str, float]:
    text = [e for e in entries if e["part"] in TEXT_PARTS]
    shares: dict[str, float] = {}
    for e in text:
        shares[e["lang"]] = shares.get(e["lang"], 0) + 1
    return {lang: n / len(text) for lang, n in shares.items()} if text else {}
