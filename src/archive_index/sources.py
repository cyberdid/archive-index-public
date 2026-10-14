"""Discover and parse sources: markdown, PDF, Office, git history, agent sessions, run logs."""

from __future__ import annotations

import fnmatch
import json
import os
import re
import subprocess
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterator

from .chunking import split_markdown_sections

SKIP_DIRS = {
    ".git", "node_modules", ".venv", "venv", "__pycache__", "site-packages",
    ".pytest_cache", ".mypy_cache", "target", "build", "dist", ".cache",
}
# Files that must never enter the index, whatever their folder (plan A3; DigUp does the same).
SENSITIVE_NAME_PATTERNS = (
    "*.pem", "*.key", "*.p12", "*.pfx", "*.kdbx", "*.gpg", "*.ppk", "id_rsa*", "id_ed25519*",
    "id_ecdsa*", "id_dsa*", ".env", ".env.*", "*.env", "credentials*", "*credential*", "*secret*",
    "*password*", "*passwd*", "*token*", "*api_key*", "*apikey*", "*.keychain*",
)
SENSITIVE_DIRS = {".ssh", ".gnupg", ".aws", ".kube", ".password-store", ".docker", "keyrings"}
# Screenshots are excluded only from the image collection, when the option is on.
SCREENSHOT_PATTERNS = ("screenshot*", "screen shot*", "знімок*", "скріншот*")
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_DOC_BYTES = 80 * 1024 * 1024
CONTEXT_PREFIXES = (
    "<environment_context>", "<codex_internal_context", "<user_instructions",
    "<command-name>", "<command-message>", "<local-command-stdout>", "<local-command-caveat>",
)


@dataclass(frozen=True)
class Unit:
    text: str
    locator: str
    heading: str = ""
    occurred_at: str = ""
    image: str = ""  # a rendered page or picture; embedded as an image, text is only for lexical search
    audio: object = None  # mono 16 kHz samples of a sound window (videos); embedded as audio


@dataclass
class Document:
    collection: str
    path: Path
    title: str
    project: str
    units: list[Unit] = field(default_factory=list)


def project_of(path: Path, projects_root: Path) -> str:
    """First path segment under the projects root ('llm' for ~/Projects/llm/x/y)."""
    try:
        parts = path.resolve().relative_to(projects_root.resolve()).parts
    except ValueError:
        return ""
    return parts[0] if len(parts) > 1 else ""


def project_from_cwd(cwd: str, projects_root: Path) -> str:
    if not cwd:
        return ""
    try:
        parts = Path(cwd).resolve().relative_to(projects_root.resolve()).parts
    except ValueError:
        return ""
    return parts[0] if parts else ""


# discovery -----------------------------------------------------------------

def is_sensitive(name: str, dirpath: str = "") -> bool:
    """True if a file name or its folder marks it as credentials or secrets."""
    lowered = name.lower()
    if any(fnmatch.fnmatch(lowered, pattern) for pattern in SENSITIVE_NAME_PATTERNS):
        return True
    parts = {part.lower() for part in Path(dirpath).parts} if dirpath else set()
    return bool(parts & SENSITIVE_DIRS)


def walk_files(root: Path, suffixes: tuple[str, ...], max_bytes: int,
               exclude_names: tuple[str, ...] = ()) -> Iterator[Path]:
    """Yield files under root with the given suffixes, pruning SKIP_DIRS and sensitive files."""
    if not root.exists():
        return
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [
            d for d in dirnames
            if (d not in SKIP_DIRS and not d.startswith(".") or d == ".antigravity_runs")
            and d.lower() not in SENSITIVE_DIRS
        ]
        for name in filenames:
            if not name.lower().endswith(suffixes):
                continue
            if is_sensitive(name, dirpath):
                continue
            if exclude_names and any(fnmatch.fnmatch(name.lower(), p) for p in exclude_names):
                continue
            path = Path(dirpath) / name
            try:
                if path.is_symlink() or path.stat().st_size > max_bytes or path.stat().st_size == 0:
                    continue
            except OSError:
                continue
            yield path


def discover_git_repos(root: Path) -> list[Path]:
    repos: list[Path] = []
    for dirpath, dirnames, _ in os.walk(root):
        if ".git" in dirnames:
            repos.append(Path(dirpath))
            dirnames.remove(".git")
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
    return repos


def discover_markdown(root: Path) -> Iterator[Path]:
    return walk_files(root, (".md",), MAX_TEXT_BYTES)


LEGACY_TARGET = {".doc": "docx", ".rtf": "docx", ".odt": "docx", ".ppt": "pptx", ".odp": "pptx"}
CONVERT_CACHE = Path.home() / ".cache" / "archive-index" / "converted"
TEXT_SKIP_DIRS = {"output", "outputs", "build", "dist", "tmp", "temp", "logs", "htmlcov", "_build", "site"}


def discover_documents(root: Path) -> Iterator[Path]:
    return walk_files(root, (".pdf", ".docx", ".pptx", ".epub", *LEGACY_TARGET), MAX_DOC_BYTES)


def parse_epub(path: Path, projects_root: Path) -> Document:
    """EPUB is a zip of XHTML files; each content file becomes sections read by the HTML parser."""
    import zipfile

    units: list[Unit] = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if not name.lower().endswith((".xhtml", ".html", ".htm")):
                continue
            html_doc = archive.read(name).decode("utf-8", errors="replace")
            for unit in _html_units(html_doc):
                units.append(Unit(text=unit.text, locator=f"{name} · {unit.locator}", heading=unit.heading))
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


# Folder names of mirrored websites look like host names: files.ettus.com, docs.python.org, ...
_HOST_LIKE = re.compile(r"^(www\.)?[a-z0-9-]+(\.[a-z0-9-]+)*\.(com|org|net|io|dev|de|ru|ua|uk|edu|gov|info|co)$")
GENERATED_DOC_MARKERS = ("generated by doxygen", "generated by sphinx", "generated by pdoc", "generated by javadoc")
THIRD_PARTY_HTML_DIRS = {"htmlcov", "_build", "site"}  # coverage reports, built docs (plan B3)


def is_mirror_or_runtime(path: Path, root: Path | None = None) -> bool:
    """A mirrored website or a Wine runtime below root: its files are third-party assets, not the user's."""
    relative = path.relative_to(root) if root is not None and path.is_relative_to(root) else path
    parts = [p.lower() for p in relative.parts[:-1]]
    if any(p.startswith("wineprefix") for p in parts):
        return True
    return "crawl" in parts or any(_HOST_LIKE.match(p) for p in parts)


def is_third_party_html(path: Path, root: Path | None = None) -> bool:
    """Documentation that is not the user's own: mirrored sites, built or generated API docs (plan B3).

    Folder rules look only below root, so a project that sits inside a folder called "site" keeps its pages.
    """
    relative = path.relative_to(root) if root is not None and path.is_relative_to(root) else path
    parts = [p.lower() for p in relative.parts]
    if any(p in THIRD_PARTY_HTML_DIRS for p in parts):
        return True
    if is_mirror_or_runtime(path, root):
        return True
    try:
        with path.open("rb") as fh:
            head = fh.read(4096).decode("utf-8", errors="replace").lower()
    except OSError:
        return False
    return any(marker in head for marker in GENERATED_DOC_MARKERS)


def discover_html(root: Path) -> Iterator[Path]:
    for path in walk_files(root, (".html", ".htm"), MAX_TEXT_BYTES):
        if not is_third_party_html(path, root):
            yield path


_MACHINE_TOKEN = re.compile(r"[0-9a-fA-F.,:;|/+\-=]+")


def is_machine_text(text: str) -> bool:
    """True for binary-like text (NUL bytes) and for dumps where most lines are numbers or hex (plan B4)."""
    if "\x00" in text[:4096]:
        return True
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return False
    dumpy = sum(1 for line in lines if _machine_line(line))
    return dumpy / len(lines) > 0.5


def _machine_line(line: str) -> bool:
    """A line where most words are numbers or hex values (a word needs a digit to count)."""
    words = line.split()
    machine = sum(1 for w in words if any(c.isdigit() for c in w) and _MACHINE_TOKEN.fullmatch(w))
    return machine / len(words) > 0.6


_TRANSCRIPT_LINE = re.compile(r"^\[(\d+):(\d+):(\d+\.\d+) --> (\d+):(\d+):(\d+\.\d+)\]\s*(.*)$")
_SPEAKER_LABEL = re.compile(r"^(?:Speaker\s*\d+)?\s*(?:\[[^\]]*\])?\s*(?:\(\d+%\))?\s*:\s*")
_NON_SPEECH_MARKERS = ("субтитр", "subtitle", "[music]", "[музик")
_AUDIO_VARIANT_SUFFIXES = ("_enhanced", "_conservative", "_3_ПОКРАЩЕНИЙ_ШІ_48k")


def _transcript_for_audio(path: Path) -> Path | None:
    """Find a companion *_transcript.txt next to the recording or under processed_recordings."""
    stem = path.stem
    for suffix in _AUDIO_VARIANT_SUFFIXES:
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    direct = path.with_name(f"{stem}_transcript.txt")
    if direct.exists():
        return direct
    # Raw recordings under <project>/recordings/<session>/<seg>.<ext> have transcripts in processed_recordings
    if path.parent.parent.name == "recordings":
        alt = path.parent.parent.parent / "processed_recordings" / path.parent.name / f"{stem}_transcript.txt"
        if alt.exists():
            return alt
    return None


def parse_audio_transcript_cues(path: Path) -> list[tuple[float, float, str]]:
    """Return (start_s, end_s, clean_text) cues from a companion transcript if one exists."""
    txt = _transcript_for_audio(path)
    if txt is None:
        return []
    cues: list[tuple[float, float, str]] = []
    try:
        lines = txt.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []
    for line in lines:
        m = _TRANSCRIPT_LINE.match(line.strip())
        if not m:
            continue
        h1, m1, s1, h2, m2, s2, raw_text = m.groups()
        words = _SPEAKER_LABEL.sub("", raw_text.strip(), count=1).strip()
        if not words:
            continue
        lowered = words.lower()
        if any(marker in lowered for marker in _NON_SPEECH_MARKERS):
            continue
        start_s = int(h1) * 3600 + int(m1) * 60 + float(s1)
        end_s = int(h2) * 3600 + int(m2) * 60 + float(s2)
        cues.append((start_s, end_s, words))
    return cues


def discover_text(root: Path) -> Iterator[Path]:
    """Plain notes: skips output folders, very large files, binary and machine-output files (plan B4)."""
    for path in walk_files(root, (".txt",), 2 * 1024 * 1024):
        inside = {p.lower() for p in path.relative_to(root).parts}  # only folders below the root count
        if inside & TEXT_SKIP_DIRS or is_mirror_or_runtime(path, root):
            continue
        try:
            sample = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        if not is_machine_text(sample):
            yield path


def discover_notebooks(root: Path) -> Iterator[Path]:
    return walk_files(root, (".ipynb",), MAX_TEXT_BYTES)


def discover_config(root: Path) -> Iterator[Path]:
    """Opt-in collection: configuration and data files, small only (plan B6)."""
    return walk_files(root, (".yaml", ".yml", ".toml", ".ini", ".json", ".xml"), 256 * 1024)


CODE_SUFFIXES = (".py", ".c", ".h", ".cpp", ".hpp", ".ts", ".js", ".rs", ".go", ".sh")
CODE_VENDOR_DIRS = {"vendor", "third_party", "third-party", "build", "dist", "target", "_build",
                    "__pycache__", "site-packages", "node_modules", "external", "deps"}


def _is_repo_doc(path: Path) -> bool:
    """README files and markdown under docs/ of a project, which go into the code collection (plan E)."""
    lowered = path.name.lower()
    return lowered.startswith("readme") or "docs" in {p.lower() for p in path.parts[:-1]}


def discover_code(root: Path, projects: set[str]) -> Iterator[Path]:
    """Source files and repository docs of the chosen projects; vendored and generated folders are skipped."""
    for path in walk_files(root, CODE_SUFFIXES + (".md",), 1024 * 1024):
        if path.suffix.lower() == ".md" and not _is_repo_doc(path):
            continue
        if projects and project_of(path, root) not in projects:
            continue
        if {p.lower() for p in path.parts} & CODE_VENDOR_DIRS:
            continue
        yield path


def parse_code(path: Path, projects_root: Path) -> Document:
    """Python: one unit per top-level function or class. Other languages: blocks split by blank lines."""
    import ast

    source = path.read_text(encoding="utf-8", errors="replace")
    project = project_of(path, projects_root)
    units: list[Unit] = []
    if path.suffix == ".py":
        try:
            tree = ast.parse(source)
        except SyntaxError:
            tree = None
        if tree is not None:
            for node in tree.body:
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    text = ast.get_source_segment(source, node) or ""
                    if text.strip():
                        units.append(Unit(text=text, locator=f"line {node.lineno}", heading=node.name))
            if units:
                return Document("code", path, str(path.relative_to(projects_root)), project, units)
    line_no = 1
    block: list[str] = []
    block_start = 1
    for line in source.splitlines(keepends=False):
        if not line.strip():
            if block:
                units.append(Unit(text="\n".join(block), locator=f"line {block_start}"))
                block = []
            line_no += 1
            block_start = line_no
            continue
        if not block:
            block_start = line_no
        block.append(line)
        line_no += 1
    if block:
        units.append(Unit(text="\n".join(block), locator=f"line {block_start}"))
    return Document("code", path, str(path.relative_to(projects_root)), project, units)


VIDEO_SUFFIXES = (".mp4", ".mkv", ".mov", ".webm", ".avi")
VIDEO_FRAME_EVERY = 5.0      # seconds between sampled frames
VIDEO_MAX_FRAMES = 120       # per video
VIDEO_MAX_SECONDS = 3600.0   # longer soundtracks are cut here
FRAME_CACHE = Path.home() / ".cache" / "archive-index" / "frames"


def discover_video(root: Path) -> Iterator[Path]:
    return walk_files(root, VIDEO_SUFFIXES, 2 * 1024 * 1024 * 1024)


def _mmss(seconds: float) -> str:
    total = int(seconds)
    return f"{total // 60:02d}:{total % 60:02d}"


def parse_video(path: Path, projects_root: Path) -> Document:
    """Frames every few seconds (images) and the soundtrack in 20 s windows (audio). PyAV bundles ffmpeg."""
    import av
    import hashlib

    import numpy as np
    from .audio_io import window_ranges

    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    out_dir = FRAME_CACHE / digest.hexdigest()
    out_dir.mkdir(parents=True, exist_ok=True)
    units: list[Unit] = []
    with av.open(str(path)) as container:
        if container.streams.video:
            next_at = 0.0
            for frame in container.decode(video=0):
                at = float(frame.time or 0.0)
                if at + 1e-6 < next_at:
                    continue
                target = out_dir / f"frame-{int(at)}.png"
                if not target.exists():
                    frame.to_image().save(target)
                units.append(Unit(text=f"{path.stem} frame {_mmss(at)}", locator=f"frame {_mmss(at)}",
                                  image=str(target)))
                next_at = at + VIDEO_FRAME_EVERY
                if len(units) >= VIDEO_MAX_FRAMES:
                    break
    with av.open(str(path)) as container:
        if container.streams.audio:
            resampler = av.AudioResampler(format="flt", layout="mono", rate=16000)
            chunks: list[np.ndarray] = []
            limit = int(VIDEO_MAX_SECONDS * 16000)
            total = 0
            for frame in container.decode(audio=0):
                for out in resampler.resample(frame):
                    chunk = out.to_ndarray().reshape(-1)
                    chunks.append(chunk)
                    total += chunk.size
                if total >= limit:
                    break
            for out in resampler.resample(None) or []:
                chunks.append(out.to_ndarray().reshape(-1))
            samples = np.concatenate(chunks).astype(np.float32) if chunks else np.zeros(0, np.float32)
            samples = samples[:limit]
            for start, end in window_ranges(len(samples) / 16000.0, 20.0, 15.0):
                piece = samples[int(start * 16000): int(end * 16000)]
                if piece.size:
                    units.append(Unit(text=f"{path.stem} sound {start:.0f}-{end:.0f}s",
                                      locator=f"{start:.0f}-{end:.0f}s", audio=piece))
    return Document("video", path, path.stem, project_of(path, projects_root), units)


def discover_logs(root: Path) -> Iterator[Path]:
    """Opt-in collection: text logs (plan B4)."""
    return walk_files(root, (".log",), 2 * 1024 * 1024)


def convert_legacy(path: Path) -> Path:
    """Convert doc/rtf/odt/ppt/odp with LibreOffice; results are cached by content hash."""
    import hashlib
    import shutil
    import tempfile

    target = LEGACY_TARGET[path.suffix.lower()]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    out = CONVERT_CACHE / f"{digest}.{target}"
    if out.exists():
        return out
    CONVERT_CACHE.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", target, "--outdir", tmp, str(path)],
                       capture_output=True, text=True, timeout=120, check=True)
        shutil.move(str(Path(tmp) / (path.stem + "." + target)), out)
    return out


def parse_legacy(path: Path, projects_root: Path) -> Document:
    converted = convert_legacy(path)
    parser = parse_docx if converted.suffix == ".docx" else parse_pptx
    doc = parser(converted, projects_root)
    doc.path = path  # the index key stays the original file
    return doc


def parse_html(path: Path, projects_root: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    units = _html_units(text)
    title = _html_title(text) or path.stem
    return Document("docs", path, title, project_of(path, projects_root), units)


def _html_title(text: str) -> str:
    import re

    match = re.search(r"<title[^>]*>(.*?)</title>", text, re.S | re.I)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def _html_units(text: str) -> list[Unit]:
    from html.parser import HTMLParser

    class _Text(HTMLParser):
        SKIP = {"script", "style", "noscript", "svg", "nav", "footer"}

        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.sections: list[tuple[str, list[str]]] = [("", [])]
            self.skip = 0
            self.heading = None
            self.title = ""

        def handle_starttag(self, tag, attrs):
            if tag in self.SKIP:
                self.skip += 1
            if tag in ("h1", "h2", "h3"):
                self.heading = ""
                self.sections.append(("", []))
            if tag == "title":
                self.heading = ""

        def handle_endtag(self, tag):
            if tag in self.SKIP and self.skip:
                self.skip -= 1
            if tag in ("h1", "h2", "h3") and self.heading is not None:
                self.sections[-1] = (self.heading.strip(), [])
                self.heading = None
            if tag == "title" and self.heading is not None:
                self.title = self.heading.strip()
                self.heading = None

        def handle_data(self, data):
            if self.skip:
                return
            if self.heading is not None:
                self.heading += data
            elif data.strip():
                self.sections[-1][1].append(data.strip())

    parser = _Text()
    parser.feed(text)
    units = []
    for heading, parts in parser.sections:
        body = " ".join(parts).strip()
        if body:
            units.append(Unit(text=body, locator=f"section {len(units) + 1}", heading=heading))
    return units


def parse_text(path: Path, projects_root: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    units = [Unit(text=text, locator="text")]
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


def parse_notebook(path: Path, projects_root: Path) -> Document:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    units = []
    for index, cell in enumerate(data.get("cells", []), start=1):
        if cell.get("cell_type") not in ("markdown", "code"):
            continue
        source = cell.get("source", [])
        text = "".join(source) if isinstance(source, list) else str(source)
        if text.strip():  # outputs are deliberately not indexed
            units.append(Unit(text=text, locator=f"cell {index}"))
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


def parse_config(path: Path, projects_root: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    return Document("config", path, path.name, project_of(path, projects_root),
                    [Unit(text=text, locator="file")])


def parse_log(path: Path, projects_root: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    return Document("logs", path, path.name, project_of(path, projects_root),
                    [Unit(text=text, locator="file")])


TABLE_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv")
MAX_TABLE_BYTES = 50 * 1024 * 1024  # larger database dumps are skipped and reported
TABLE_WINDOW_ROWS = 30
TABLE_MAX_ROWS = 20000  # per sheet / file


def discover_tables(root: Path) -> Iterator[Path]:
    return walk_files(root, TABLE_SUFFIXES, MAX_TABLE_BYTES)


def _cells_text(row) -> str:
    return " | ".join("" if v is None else str(v).strip() for v in row).strip(" |")


TABLE_WINDOW_CHARS = 1200  # a window also ends here, so one long row cannot make a huge input


TABLE_CELL_CHARS = 200  # binary blobs and long hex dumps are cut; they do not help search


def _sheet_units(sheet_name: str, rows: list[list[str]]) -> list[Unit]:
    """Windows of rows; every window repeats the header so its numbers keep their column names."""
    rows = [[(c[:TABLE_CELL_CHARS] + "…") if len(c) > TABLE_CELL_CHARS else c for c in r] for r in rows]
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    header_text = " | ".join(c for c in rows[0] if c)
    body = [" | ".join(r).strip(" |") for r in rows[1:]]
    units: list[Unit] = []
    window: list[str] = []
    first = 2
    row_no = 1  # the header is row 1

    def flush(last_row: int) -> None:
        if window:
            text = "\n".join([f"[{header_text}]"] + window)
            units.append(Unit(text=text, locator=f"sheet {sheet_name} · rows {first}-{last_row}"))

    for row_no, line in enumerate(body, start=2):
        if window and (len(window) >= TABLE_WINDOW_ROWS or
                       sum(len(x) for x in window) + len(line) > TABLE_WINDOW_CHARS):
            flush(row_no - 1)
            window = []
            first = row_no
        window.append(line)
    flush(row_no)
    if not units and rows:
        units.append(Unit(text=f"[{header_text}]", locator=f"sheet {sheet_name} · rows 1-1"))
    return units


def _read_csv(path: Path) -> list[list[str]]:
    import csv

    csv.field_size_limit(50 * 1024 * 1024)  # database dumps can hold large fields

    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = []
    for i, row in enumerate(csv.reader(text.splitlines(), dialect)):
        if i >= TABLE_MAX_ROWS:
            break
        rows.append(row)
    return rows


def _read_workbook(path: Path) -> list[tuple[str, list[list[str]]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in workbook.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= TABLE_MAX_ROWS:
                    break
                rows.append(["" if v is None else str(v) for v in row])
            sheets.append((ws.title, rows))
    finally:
        workbook.close()
    return sheets


def parse_table(path: Path, projects_root: Path) -> Document:
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        sheets = [(path.stem, _read_csv(path))]
    elif suffix in (".xls", ".ods"):
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                capture_output=True, text=True, timeout=120, check=True,
            )
            converted = Path(tmp) / (path.stem + ".xlsx")
            sheets = _read_workbook(converted)
    else:
        sheets = _read_workbook(path)
    units: list[Unit] = []
    for name, rows in sheets:
        units.extend(_sheet_units(name, rows))
    return Document("tables", path, path.stem, project_of(path, projects_root), units)


def discover_run_logs(root: Path) -> Iterator[Path]:
    for path in walk_files(root, (".jsonl",), MAX_TEXT_BYTES * 20):
        if path.parent.name == ".antigravity_runs":
            yield path


IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".webp")
AUDIO_SUFFIXES = (".wav", ".mp3", ".opus", ".ogg", ".flac", ".m4a")


def discover_images(root: Path, skip_screenshots: bool = False) -> Iterator[Path]:
    excluded = SCREENSHOT_PATTERNS if skip_screenshots else ()
    for path in walk_files(root, IMAGE_SUFFIXES, 15 * 1024 * 1024, exclude_names=excluded):
        if not is_mirror_or_runtime(path, root):  # images of mirrored sites and Wine runtimes are not notes
            yield path


def discover_audio(root: Path) -> Iterator[Path]:
    return walk_files(root, AUDIO_SUFFIXES, 300 * 1024 * 1024)


def discover_claude_sessions(home: Path) -> list[Path]:
    return sorted((home / ".claude" / "projects").glob("*/*.jsonl"))


def discover_codex_sessions(home: Path) -> list[Path]:
    return sorted((home / ".codex" / "sessions").rglob("rollout-*.jsonl"))


# parsers -------------------------------------------------------------------

def parse_markdown(path: Path, projects_root: Path) -> Document:
    text = path.read_text(encoding="utf-8", errors="replace")
    units = [
        Unit(text=body, locator=f"section {i + 1}", heading=heading)
        for i, (heading, body) in enumerate(split_markdown_sections(text))
    ]
    title = next((h for h, _ in split_markdown_sections(text) if h), path.stem)
    return Document("wiki", path, title, project_of(path, projects_root), units)


PAGE_CACHE = Path.home() / ".cache" / "archive-index" / "pages"
SCAN_MIN_CHARS = 50  # a page with less text than this is treated as a scan or picture


def _render_page(path: Path, page: int) -> Path | None:
    """Render one PDF page to PNG (cached by file hash); None if rendering fails."""
    import hashlib

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    target_dir = PAGE_CACHE / digest
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"page-{page}.png"
    if target.exists():
        return target
    proc = subprocess.run(
        ["pdftoppm", "-png", "-r", "100", "-f", str(page), "-l", str(page), "-singlefile",
         str(path), str(target.with_suffix(""))],
        capture_output=True, text=True, timeout=120,
    )
    return target if proc.returncode == 0 and target.exists() else None


def parse_pdf(path: Path, projects_root: Path, max_page: int | None = None) -> Document:
    """max_page limits rendering of image-only pages (staged indexing renders only the current pass)."""
    proc = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", str(path), "-"],
        capture_output=True, text=True, timeout=180, check=True,
    )
    units: list[Unit] = []
    for i, page in enumerate(proc.stdout.split("\f")):
        text = page.strip()
        if len(text) >= SCAN_MIN_CHARS:
            units.append(Unit(text=text, locator=f"page {i + 1}"))
            continue
        if max_page is not None and i >= max_page:
            break
        rendered = _render_page(path, i + 1) if i < 2000 else None  # cap work on huge scans
        if rendered is not None:
            units.append(Unit(text=f"{path.stem} page {i + 1} image", locator=f"page {i + 1} (image)",
                              image=str(rendered)))
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


def parse_docx(path: Path, projects_root: Path) -> Document:
    import docx  # python-docx

    paragraphs = [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]
    units = [
        Unit(text="\n".join(paragraphs[i : i + 40]), locator=f"paragraphs {i + 1}")
        for i in range(0, len(paragraphs), 40)
    ]
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


def parse_pptx(path: Path, projects_root: Path) -> Document:
    from pptx import Presentation

    units: list[Unit] = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        body = "\n".join(t for t in texts if t.strip())
        if body.strip():
            units.append(Unit(text=body, locator=f"slide {index}"))
    return Document("docs", path, path.stem, project_of(path, projects_root), units)


def parse_git_history(repo: Path, projects_root: Path, limit: int = 500) -> Document:
    fmt = "%H%x1f%ad%x1f%s%x1f%b%x1e"
    proc = subprocess.run(
        ["git", "-C", str(repo), "log", f"-n{limit}", f"--format={fmt}", "--date=short"],
        capture_output=True, text=True, timeout=120, check=False,
    )
    units: list[Unit] = []
    for record in proc.stdout.split("\x1e"):
        parts = record.strip("\n").split("\x1f")
        if len(parts) < 3 or not parts[0]:
            continue
        sha, date, subject = parts[0], parts[1], parts[2]
        body = parts[3].strip() if len(parts) > 3 else ""
        text = subject if not body else f"{subject}\n\n{body}"
        units.append(Unit(text=text, locator=f"commit {sha[:10]}", occurred_at=date))
    title = f"git history: {repo.name}"
    return Document("git", repo / ".git", title, project_of(repo, projects_root), units)


def _text_of_content(content: object) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        return "\n".join(
            item.get("text", "") for item in content
            if isinstance(item, dict) and item.get("type") in ("text", "input_text", "output_text")
        )
    return ""


def _turns_to_units(turns: list[dict]) -> list[Unit]:
    units = []
    for turn in turns:
        body = "\n\n".join(
            f"{speaker}: {text}" for speaker, text in turn["messages"] if text.strip()
        )
        if body:
            units.append(Unit(text=body, locator=turn["locator"], occurred_at=turn["at"]))
    return units


def _is_context_noise(text: str) -> bool:
    stripped = text.lstrip()
    return stripped.startswith(CONTEXT_PREFIXES)


def _iter_complete_lines(path: Path, start_offset: int, start_line: int):
    """Yield (line_number, byte_offset, raw_bytes) for complete lines from start_offset on.

    A last line without a trailing newline is still being written; it is not consumed.
    """
    with path.open("rb") as fh:
        fh.seek(start_offset)
        offset, lineno = start_offset, start_line
        for raw in fh:
            if not raw.endswith(b"\n"):
                break
            yield lineno, offset, raw
            offset += len(raw)
            lineno += 1


def _build_turns(path, projects_root, start_offset, start_line, is_record, extract):
    """Shared session walker.

    is_record(record) -> (kind, text, timestamp) for user/assistant messages, or None.
    Returns (turns, project, session_id, first_user, resume_offset, resume_line).
    The resume point is the start of the last user turn, so that turn is re-read next time.
    """
    turns: list[dict] = []
    project = ""
    session_id = path.stem
    first_user = ""
    resume_offset, resume_line = start_offset, start_line
    for lineno, offset, raw in _iter_complete_lines(path, start_offset, start_line):
        resume_offset, resume_line = offset + len(raw), lineno + 1
        try:
            record = json.loads(raw.decode("utf-8", errors="replace"))
        except json.JSONDecodeError:
            continue
        if not isinstance(record, dict):
            continue
        meta_project, meta_session = extract(record)
        if meta_project and not project:
            project = meta_project
        if meta_session:
            session_id = meta_session
        parsed = is_record(record)
        if parsed is None:
            continue
        kind, text, at = parsed
        if kind == "assistant":
            if text and turns:
                turns[-1]["messages"].append(("ASSISTANT", text))
            continue
        if not text or _is_context_noise(text):
            continue
        if not first_user:
            first_user = text.splitlines()[0][:80]
        turns.append({"messages": [("USER", text)], "locator": f"line {lineno}", "at": at,
                      "offset": offset, "line": lineno})
    if turns:
        resume_offset, resume_line = turns[-1]["offset"], turns[-1]["line"]
    return turns, project, session_id, first_user, resume_offset, resume_line


def parse_claude_session(path: Path, projects_root: Path, start_offset: int = 0,
                         start_line: int = 1):
    """Return (Document, resume_offset, resume_line) for a Claude Code transcript."""
    def is_record(record):
        kind = record.get("type")
        if kind not in ("user", "assistant"):
            return None
        message = record.get("message") or {}
        return kind, _text_of_content(message.get("content")).strip(), record.get("timestamp", "")

    def extract(record):
        cwd = record.get("cwd", "")
        return (project_from_cwd(cwd, projects_root) if cwd else ""), record.get("sessionId", "")

    turns, project, session_id, first_user, r_off, r_line = _build_turns(
        path, projects_root, start_offset, start_line, is_record, extract)
    title = f"claude session {session_id[:8]}: {first_user}".strip(": ")
    doc = Document("sessions", path, title, project, _turns_to_units(turns))
    return doc, r_off, r_line


def parse_codex_session(path: Path, projects_root: Path, start_offset: int = 0,
                        start_line: int = 1):
    """Return (Document, resume_offset, resume_line) for a Codex rollout transcript."""
    def is_record(record):
        if record.get("type") != "response_item":
            return None
        payload = record.get("payload") or {}
        if not isinstance(payload, dict) or payload.get("type") != "message":
            return None
        role = payload.get("role")
        if role not in ("user", "assistant"):
            return None
        kind = "user" if role == "user" else "assistant"
        return kind, _text_of_content(payload.get("content")).strip(), record.get("timestamp", "")

    def extract(record):
        payload = record.get("payload") or {}
        if not isinstance(payload, dict):
            return "", ""
        if record.get("type") in ("session_meta", "turn_context"):
            cwd = payload.get("cwd", "")
            return (project_from_cwd(cwd, projects_root) if cwd else ""), payload.get("id", "")
        return "", ""

    turns, project, session_id, first_user, r_off, r_line = _build_turns(
        path, projects_root, start_offset, start_line, is_record, extract)
    title = f"codex session {session_id[:8]}: {first_user}".strip(": ")
    doc = Document("sessions", path, title, project, _turns_to_units(turns))
    return doc, r_off, r_line


RUN_EVENTS = {
    "attempt_result", "subtask_failed", "subtask_complete", "replan",
    "subtask_paused", "subtask_blocked", "verification_failed", "footer",
}


def parse_run_log(path: Path, projects_root: Path) -> Document:
    units: list[Unit] = []
    project = project_of(path, projects_root)
    with path.open(encoding="utf-8", errors="replace") as fh:
        for lineno, line in enumerate(fh, start=1):
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue
            event = record.get("event_type") or record.get("kind")
            if event not in RUN_EVENTS and record.get("kind") != "footer":
                continue
            payload = record.get("payload") or record
            pieces = [str(event)]
            for key in ("subtask", "subtask_id", "error", "reason", "summary", "title"):
                value = payload.get(key) if isinstance(payload, dict) else None
                if value:
                    pieces.append(f"{key}: {str(value)[:600]}")
            units.append(Unit(text="\n".join(pieces), locator=f"line {lineno}"))
    return Document("run_logs", path, f"run log {path.stem}", project, units)


def first_line_date(value: str) -> str:
    """Normalize an ISO-ish timestamp to YYYY-MM-DD when possible."""
    match = re.match(r"(\d{4}-\d{2}-\d{2})", value or "")
    if match:
        try:
            datetime.strptime(match.group(1), "%Y-%m-%d")
            return match.group(1)
        except ValueError:
            return ""
    return ""
